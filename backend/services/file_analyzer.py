"""
Phân tích sơ bộ file upload — trả về summary + tool suggestions
Hỗ trợ: Excel, CSV, PDF, Word, Image, Text
"""

import os
import io
from typing import Optional


def analyze_file(path: str, filename: str) -> dict:
    ext = os.path.splitext(filename.lower())[1]

    if ext in (".xlsx", ".xls"):
        return _analyze_excel(path, filename)
    elif ext == ".csv":
        return _analyze_csv(path, filename)
    elif ext == ".pdf":
        return _analyze_pdf(path, filename)
    elif ext in (".docx", ".doc"):
        return _analyze_word(path, filename)
    elif ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"):
        return _analyze_image(path, filename)
    elif ext in (".txt", ".md"):
        return _analyze_text(path, filename)
    else:
        return {
            "type": "unknown",
            "filename": filename,
            "summary": f"Loại file {ext} — không hỗ trợ phân tích tự động.",
            "tool_suggestions": [],
            "meta": {},
        }


# ── Excel ─────────────────────────────────────────────────────────────────

def _analyze_excel(path: str, filename: str) -> dict:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets_info = []
        tool_suggestions = []

        for sname in wb.sheetnames[:5]:
            ws = wb[sname]
            rows = list(ws.iter_rows(min_row=1, max_row=6, values_only=True))
            if not rows:
                continue
            headers = [str(c) if c is not None else "" for c in rows[0]]
            preview = [
                {headers[j]: str(v) if v is not None else ""
                 for j, v in enumerate(row) if j < len(headers)}
                for row in rows[1:4]
            ]
            max_row = ws.max_row or 0
            sheets_info.append({
                "name": sname,
                "rows": max_row,
                "columns": headers,
                "preview": preview,
            })

            # Detect tool suggestions based on column names
            col_lower = " ".join(headers).lower()
            if any(kw in col_lower for kw in ["budget", "actual", "ngân sách", "thực tế", "variance", "chênh lệch", "kế hoạch"]):
                if "analyze_variance" not in tool_suggestions:
                    tool_suggestions.append("analyze_variance")
            if any(kw in col_lower for kw in ["bank", "ngân hàng", "sao kê", "transaction", "debit", "credit", "số dư"]):
                if "bank_reconciliation" not in tool_suggestions:
                    tool_suggestions.append("bank_reconciliation")

        wb.close()

        summary_lines = [f"**{filename}** — Excel workbook với {len(wb.sheetnames)} sheet(s)"]
        for s in sheets_info:
            summary_lines.append(
                f"- **{s['name']}**: {s['rows']} dòng × {len(s['columns'])} cột"
                + (f" — Cột: {', '.join(s['columns'][:6])}" if s['columns'] else "")
            )

        return {
            "type": "excel",
            "filename": filename,
            "summary": "\n".join(summary_lines),
            "tool_suggestions": tool_suggestions,
            "meta": {"sheets": sheets_info},
        }
    except Exception as e:
        return {"type": "excel", "filename": filename, "summary": f"Lỗi đọc Excel: {e}", "tool_suggestions": [], "meta": {}}


# ── CSV ───────────────────────────────────────────────────────────────────

def _analyze_csv(path: str, filename: str) -> dict:
    try:
        import csv
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.reader(f)
            rows = [r for r, _ in zip(reader, range(6))]
        headers = rows[0] if rows else []
        row_count = sum(1 for _ in open(path, encoding="utf-8-sig", errors="replace")) - 1

        col_lower = " ".join(headers).lower()
        suggestions = []
        if any(kw in col_lower for kw in ["budget", "actual", "ngân sách", "thực tế", "variance"]):
            suggestions.append("analyze_variance")
        if any(kw in col_lower for kw in ["bank", "ngân hàng", "sao kê", "debit", "credit"]):
            suggestions.append("bank_reconciliation")

        return {
            "type": "csv",
            "filename": filename,
            "summary": f"**{filename}** — {row_count} dòng × {len(headers)} cột\nCột: {', '.join(headers[:8])}",
            "tool_suggestions": suggestions,
            "meta": {"headers": headers, "row_count": row_count},
        }
    except Exception as e:
        return {"type": "csv", "filename": filename, "summary": f"Lỗi đọc CSV: {e}", "tool_suggestions": [], "meta": {}}


# ── PDF ───────────────────────────────────────────────────────────────────

def _analyze_pdf(path: str, filename: str) -> dict:
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        pages = len(reader.pages)
        text = ""
        for p in reader.pages[:3]:
            text += p.extract_text() or ""
        excerpt = text[:500].replace("\n", " ").strip()
        return {
            "type": "pdf",
            "filename": filename,
            "summary": f"**{filename}** — PDF {pages} trang\nNội dung đầu: {excerpt}{'…' if len(text) > 500 else ''}",
            "tool_suggestions": [],
            "meta": {"pages": pages, "excerpt": excerpt},
        }
    except Exception as e:
        return {"type": "pdf", "filename": filename, "summary": f"Lỗi đọc PDF: {e}", "tool_suggestions": [], "meta": {}}


# ── Word ──────────────────────────────────────────────────────────────────

def _analyze_word(path: str, filename: str) -> dict:
    try:
        from docx import Document
        doc = Document(path)
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        excerpt = " ".join(paras[:6])[:500]
        return {
            "type": "word",
            "filename": filename,
            "summary": f"**{filename}** — Word document {len(doc.paragraphs)} đoạn\nNội dung đầu: {excerpt}{'…' if len(excerpt) == 500 else ''}",
            "tool_suggestions": [],
            "meta": {"paragraphs": len(doc.paragraphs), "excerpt": excerpt},
        }
    except Exception as e:
        return {"type": "word", "filename": filename, "summary": f"Lỗi đọc Word: {e}", "tool_suggestions": [], "meta": {}}


# ── Image ─────────────────────────────────────────────────────────────────

def _analyze_image(path: str, filename: str) -> dict:
    size_kb = os.path.getsize(path) // 1024
    try:
        from PIL import Image
        with Image.open(path) as img:
            w, h = img.size
            mode = img.mode
        return {
            "type": "image",
            "filename": filename,
            "summary": f"**{filename}** — Hình ảnh {w}×{h}px, {mode}, {size_kb}KB",
            "tool_suggestions": [],
            "meta": {"width": w, "height": h, "size_kb": size_kb},
        }
    except Exception:
        return {
            "type": "image",
            "filename": filename,
            "summary": f"**{filename}** — Hình ảnh {size_kb}KB",
            "tool_suggestions": [],
            "meta": {"size_kb": size_kb},
        }


# ── Text / Markdown ───────────────────────────────────────────────────────

def _analyze_text(path: str, filename: str) -> dict:
    try:
        with open(path, encoding="utf-8", errors="replace") as f:
            content = f.read()
        lines = len(content.splitlines())
        excerpt = content[:400].strip()
        return {
            "type": "text",
            "filename": filename,
            "summary": f"**{filename}** — {lines} dòng\nNội dung: {excerpt}{'…' if len(content) > 400 else ''}",
            "tool_suggestions": [],
            "meta": {"lines": lines},
        }
    except Exception as e:
        return {"type": "text", "filename": filename, "summary": f"Lỗi đọc file: {e}", "tool_suggestions": [], "meta": {}}
